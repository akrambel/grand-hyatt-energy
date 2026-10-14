#!/usr/bin/env python3
"""
Generates the test fixtures used by tests/*.test.ts.

  1. tests/fixtures/template-cells.json  – labels / dates / constants of the real
     workbook template (formulas excluded) so the importer & exporter can be tested
     without any spreadsheet library.
  2. tests/fixtures/sample-input.json    – deterministic TEST DATA (not real readings).
  3. tests/fixtures/sample-expected.json – what the workbook's OWN formulas compute for
     that test data. The workbook is filled with the sample data, recalculated with
     LibreOffice, and the cached results are read back. tests/calc-vs-workbook.test.ts
     then asserts that the TypeScript engine produces the same numbers.

Requires: python3, openpyxl, LibreOffice (soffice) and the recalc helper script.

  python3 scripts/make_fixtures.py public/templates/Main_Consumption_Opening_Closing.xlsx \
          --recalc /path/to/recalc.py
"""
import argparse
import datetime as dt
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures"


def cell_value(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return {"$date": v.strftime("%Y-%m-%d")}
    return v


def dump_template(wb):
    out = {}
    for name in ("Daily Consumption", "Utility Budget & Actual", "LPG Dashboard"):
        ws = wb[name]
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or (isinstance(v, str) and (v.startswith("=") or v.strip() == "")):
                    continue
                cells[f"{c.row},{c.column}"] = cell_value(v)
        out[name] = cells
    return out


class Lcg:
    def __init__(self, seed):
        self.s = seed

    def next(self):
        self.s = (self.s * 1103515245 + 12345) & 0x7FFFFFFF
        return self.s / 0x7FFFFFFF


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("--recalc", required=True)
    args = ap.parse_args()

    FIX.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(args.template)
    (FIX / "template-cells.json").write_text(json.dumps(dump_template(wb), ensure_ascii=False, separators=(",", ":")))

    ws = wb["Daily Consumption"]
    sections = {
        "energy": (7, 16, 17, 1500.0, 40.0),
        "water": (22, 31, 32, 800.0, 3.0),
        "gas": (37, 45, 46, 20000.0, 120.0),
        "chilled": (51, 70, 106, 5000.0, 12.0),  # only the first 20 BTU meters get test data
    }
    start = dt.date(2026, 10, 1)
    days = 20  # 1..20 October 2026
    rnd = Lcg(2026)
    readings = []
    # climate / occupancy rows
    metrics = []
    for d in range(days):
        date = start + dt.timedelta(days=d)
        col = 2 + d
        t = round(31 + 4 * rnd.next(), 1)
        occ = round(0.55 + 0.35 * rnd.next(), 3)
        rooms = int(occ * 200)
        guests = int(rooms * 2.1)
        # leave occupancy blank on day 5 and temperature blank on day 7 to test missing data
        row = {"date": date.isoformat()}
        if d != 6:
            ws.cell(110, col, t)
            row["avg_temp_c"] = t
        if d != 4:
            ws.cell(112, col, occ)
            ws.cell(113, col, rooms)
            ws.cell(114, col, guests)
            row.update(occupancy_rate=occ, occupied_rooms=rooms, guests=guests)
        metrics.append(row)

    for util, (r0, r1, _t, base, step) in sections.items():
        for r in range(r0, r1 + 1):
            name = ws.cell(r, 1).value
            level = base * (1 + 0.1 * (r - r0))
            for d in range(days):
                # test gaps: one energy meter has no reading on day 3, one water meter starts on day 4
                if util == "energy" and r == r0 + 2 and d == 2:
                    continue
                if util == "water" and r == r0 + 1 and d < 3:
                    level_start = True
                    continue
                level += step * (0.6 + 0.8 * rnd.next())
                level = round(level, 2)
                ws.cell(r, 2 + d, level)
                readings.append({"utility": util, "row": r, "meter": name.strip(),
                                 "date": (start + dt.timedelta(days=d)).isoformat(), "value": level})

    # budget inputs (TEST DATA)
    wsb = wb["Utility Budget & Actual"]
    budget = {}
    for i, key in enumerate(["rooms", "occ", "guests", None, "gas", "electricity", "water"]):
        if key is None:
            continue
    rows = {"rooms_sold": 7, "occupancy_rate": 8, "guests": 9, "gas": 11, "electricity": 12, "water": 13}
    vals = {"rooms_sold": 4000, "occupancy_rate": 0.7, "guests": 8400, "gas": 90000, "electricity": 400000, "water": 9000}
    for k, r in rows.items():
        for c in range(2, 14):
            wsb.cell(r, c, vals[k] * (1 + 0.01 * c))
        budget[k] = [round(vals[k] * (1 + 0.01 * c), 4) for c in range(2, 14)]

    tmp = Path(tempfile.mkdtemp())
    f = tmp / "sample.xlsx"
    wb.save(f)
    subprocess.run([sys.executable, args.recalc, str(f)], check=True, cwd=str(Path(args.recalc).parent))

    res = openpyxl.load_workbook(f, data_only=True)
    dc = res["Daily Consumption"]

    def num(v):
        return None if v in (None, "") else round(float(v), 6)

    expected = {"daily": {}, "month_oct": {}, "summary_oct": {}, "budget_actual": {}}
    for util, (_a, _b, t, _base, _s) in sections.items():
        expected["daily"][util] = {(start + dt.timedelta(days=d)).isoformat(): num(dc.cell(t, 2 + d).value)
                                   for d in range(days)}
    octs = res["Oct 2026"]
    expected["month_oct"] = {
        "energy": num(octs["B39"].value), "water": num(octs["C39"].value),
        "gas": num(octs["D39"].value), "chilled": num(octs["E39"].value),
    }
    ms = res["Monthly Summary"]
    expected["summary_oct"] = {"energy": num(ms["B8"].value), "water": num(ms["C8"].value),
                               "gas": num(ms["D8"].value), "chilled": num(ms["E8"].value),
                               "all": num(ms["F8"].value)}
    ba = res["Utility Budget & Actual"]
    expected["budget_actual"] = {
        "rooms_sold": num(ba["B18"].value), "occupancy_avg": num(ba["B19"].value),
        "guests": num(ba["B20"].value), "gas": num(ba["B22"].value),
        "electricity": num(ba["B23"].value), "water": num(ba["B24"].value), "total": num(ba["B25"].value),
        "budget_total_oct": num(ba["B14"].value),
    }

    (FIX / "sample-input.json").write_text(json.dumps(
        {"_comment": "TEST DATA generated by scripts/make_fixtures.py - not real meter readings",
         "readings": readings, "metrics": metrics, "budget": budget}, indent=0, separators=(",", ":")))
    (FIX / "sample-expected.json").write_text(json.dumps(
        {"_comment": "Values computed by the workbook's own formulas (LibreOffice recalculation)",
         **expected}, indent=1))
    shutil.copy(f, tmp / "keep.xlsx")
    print("fixtures written to", FIX)


if __name__ == "__main__":
    main()
